"""XLSX adapter for Crusoe's normalized workbook."""

from __future__ import annotations

import hashlib
import re
import tempfile
from pathlib import Path
import os
from datetime import datetime, timezone

from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.output.model import OutputResult, Workbook

EXCEL_CELL_LIMIT = 32_767


def xlsx_destination(topic: str, config: dict, config_dir: str | Path) -> Path:
    output = config.get("output", {}) if isinstance(config.get("output", {}), dict) else {}
    xlsx = output.get("xlsx", {}) if isinstance(output.get("xlsx", {}), dict) else {}
    directory = Path(xlsx.get("directory", "data/outputs"))
    if not directory.is_absolute():
        directory = Path(config_dir) / directory
    slug = re.sub(r"[^a-z0-9]+", "-", topic.lower()).strip("-") or "research"
    slug = slug[:60]
    digest = hashlib.sha256(topic.encode("utf-8")).hexdigest()[:8]
    return directory / f"{slug}-{digest}.xlsx"


def _safe_text(value: object) -> str:
    text = "" if value is None else str(value)
    if len(text) > EXCEL_CELL_LIMIT:
        suffix = "… [truncated: Excel cell limit reached]"
        return text[: EXCEL_CELL_LIMIT - len(suffix)] + suffix
    # Formula-looking values are assigned as strings below, never formulas.
    return "".join(ch for ch in text if ord(ch) >= 32 or ch in "\t\n\r")


def write_xlsx(workbook: Workbook, path: str | Path) -> str:
    """Atomically write an XLSX workbook and return its absolute path."""
    target = Path(path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=target.parent, suffix=".tmp", delete=False) as temp:
        temp_path = Path(temp.name)
    try:
        output = OpenpyxlWorkbook()
        output.remove(output.active)
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for sheet in workbook.sheets:
            target_sheet = output.create_sheet(sheet.name)
            target_sheet.append(list(sheet.headers))
            for row in sheet.rows:
                target_sheet.append([_safe_text(cell.value) if isinstance(cell.value, str) else cell.value for cell in row])
            # openpyxl otherwise interprets a leading '=' as a formula.  Mark
            # every projected string as text so untrusted paper content stays
            # literal while preserving its visible value.
            for row in target_sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            for column, header in enumerate(sheet.headers, 1):
                header_cell = target_sheet.cell(1, column)
                header_cell.fill = header_fill
                header_cell.font = Font(bold=True, color="FFFFFF")
                header_cell.alignment = Alignment(wrap_text=True, vertical="top")
                target_sheet.column_dimensions[get_column_letter(column)].width = min(55, max(12, len(header) + 2))
            for row in target_sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            for row_number, row in enumerate(sheet.rows, 2):
                for column, cell in enumerate(row, 1):
                    if cell.hyperlink:
                        target_sheet.cell(row_number, column).hyperlink = cell.hyperlink
                        target_sheet.cell(row_number, column).style = "Hyperlink"
            target_sheet.freeze_panes = "A2"
            target_sheet.auto_filter.ref = f"A1:{get_column_letter(max(1, len(sheet.headers)))}{max(1, target_sheet.max_row)}"
        output.save(temp_path)
        os.replace(temp_path, target)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    return str(target)


class XlsxWriter:
    """XLSX implementation of the shared spreadsheet-writer contract."""

    def __init__(self, destination: str | Path) -> None:
        self.destination = str(Path(destination).resolve())

    def ensure_destination(
        self, workbook: Workbook, destination_id: str | None = None
    ) -> str:
        return destination_id or self.destination

    def write(
        self, workbook: Workbook, destination_id: str, payload_sha256: str
    ) -> OutputResult:
        location = write_xlsx(workbook, destination_id)
        return OutputResult(
            status="success",
            location=location,
            destination_id=destination_id,
            payload_sha256=payload_sha256,
            completed_at=datetime.now(timezone.utc).isoformat(),
        )
