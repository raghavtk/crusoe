from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.core.state import PipelineState
from src.output import (
    GoogleSheetsWriter,
    payload_fingerprint,
    project_state,
    validate_output_config,
    write_xlsx,
    xlsx_destination,
)


def _state() -> PipelineState:
    return PipelineState(
        topic="Evidence synthesis / 2026",
        papers_curated=[{
            "paperId": "p-1", "title": "=Literal formula", "url": "https://example.test/paper",
            "authors": [{"name": "Ada Lovelace"}, {"name": "Grace Hopper"}], "year": 2025,
            "citationCount": 42, "fieldsOfStudy": ["Computer Science", "Security"],
            "assessment_status": "success", "methodology": "experiment",
            "contribution_type": "method", "relevance_score": 5, "confidence_score": 0.9,
            "reading_priority_score": 9, "reading_priority": "high", "relevance_rationale": "Useful\nresult",
            "one_line_summary": "A summary", "abstract": "Abstract ✓",
        }],
        synthesis={
            "summary_paragraph": "Summary", "key_themes": ["Legacy theme"],
            "suggested_reading_order": [{"paperId": "p-1", "title": "Paper", "reason": "Start here"}],
            "landscape": {
                "themes": [{"name": "Theme", "explanation": "Evidence", "supporting_paper_ids": ["p-1"], "confidence": 0.8}],
                "gaps": [{"name": "Gap", "explanation": "Unknown", "supporting_paper_ids": ["p-1"], "confidence": 0.6}],
                "future_work": [{"recommendation": "Test", "rationale": "Reason", "supporting_paper_ids": ["p-1"], "confidence": 0.7}],
                "methodology_patterns": [{"methodology": "experiment", "observation": "common", "representative_paper_ids": ["p-1"]}],
                "disagreements": [{"question": "Does it work?", "positions": [{"position": "yes", "supporting_paper_ids": ["p-1"]}], "interpretation": "mixed"}],
                "shared_limitations": [{"limitation": "Small sample", "supporting_paper_ids": ["p-1"]}],
            },
        },
    )


def test_projector_is_deterministic_and_normalizes_papers() -> None:
    workbook = project_state(_state())
    assert [sheet.name for sheet in workbook.sheets] == ["Summary", "Papers", "Themes", "Gaps", "Future Work", "Methods", "Disagreements", "Reading Order"]
    papers = workbook.sheets[1]
    assert papers.rows[0][0].value == "p-1"
    assert papers.rows[0][3].value == "Ada Lovelace; Grace Hopper"
    assert papers.rows[0][4].value == "Computer Science; Security"
    assert papers.rows[0][2].hyperlink == "https://example.test/paper"
    assert workbook.sheets[2].rows[0][3].value == 0.8
    assert workbook.sheets[3].rows[1][0].value == "Shared limitation"
    assert payload_fingerprint(workbook) == payload_fingerprint(project_state(_state()))


def test_xlsx_writer_round_trips_format_links_and_literal_formula(tmp_path: Path) -> None:
    workbook = project_state(_state())
    output = xlsx_destination(_state().topic, {}, tmp_path)
    assert output.name.startswith("evidence-synthesis-2026-")
    write_xlsx(workbook, output)
    loaded = load_workbook(output)
    assert loaded.sheetnames == [sheet.name for sheet in workbook.sheets]
    papers = loaded["Papers"]
    assert papers.freeze_panes == "A2"
    assert papers["B2"].value == "=Literal formula"
    assert papers["B2"].data_type == "s"
    assert papers["C2"].hyperlink.target == "https://example.test/paper"
    assert papers["G2"].value == 42
    assert papers.auto_filter.ref == "A1:Q2"


def test_xlsx_marks_cell_limit_truncation(tmp_path: Path) -> None:
    state = _state()
    state.papers_curated[0]["abstract"] = "x" * 40_000
    output = tmp_path / "book.xlsx"
    write_xlsx(project_state(state), output)
    value = load_workbook(output)["Papers"]["Q2"].value
    assert len(value) == 32_767
    assert value.endswith("[truncated: Excel cell limit reached]")


def test_google_writer_preserves_unmanaged_tabs_and_writes_raw() -> None:
    class Values:
        def __init__(self) -> None: self.calls = []
        def clear(self, **kwargs): self.calls.append(("clear", kwargs)); return self
        def update(self, **kwargs): self.calls.append(("update", kwargs)); return self
        def execute(self): return {}
    class Sheets:
        def __init__(self, values): self.values_api = values; self.requests = []
        def get(self, **kwargs): self.requests.append(("get", kwargs)); return self
        def batchUpdate(self, **kwargs): self.requests.append(("batch", kwargs)); return self
        def values(self): return self.values_api
        def execute(self): return {"sheets": [{"properties": {"title": name, "sheetId": index}} for index, name in enumerate(["Notes", "Summary", "Papers", "Themes", "Gaps", "Future Work", "Methods", "Disagreements", "Reading Order"])]}
    class Service:
        def __init__(self): self.values_api = Values(); self.api = Sheets(self.values_api)
        def spreadsheets(self): return self.api
    google_state = _state()
    google_state.papers_curated[0]["abstract"] = "x" * 60_000
    service = Service()
    result = GoogleSheetsWriter(service=service).write(
        project_state(google_state), "sheet-1", "payload-hash"
    )
    assert result.location is not None and result.location.endswith("sheet-1")
    assert result.payload_sha256 == "payload-hash"
    assert all("Notes" not in call[1].get("range", "") for call in service.values_api.calls)
    updates = [call[1] for call in service.values_api.calls if call[0] == "update"]
    assert updates and all(call["valueInputOption"] == "RAW" for call in updates)
    papers_update = next(call for call in updates if call["range"] == "'Papers'!A1")
    abstract = papers_update["body"]["values"][1][16]
    assert len(abstract) == 50_000
    assert abstract.endswith("[truncated: Google Sheets cell limit reached]")
    batch_requests = [
        request
        for kind, call in service.api.requests if kind == "batch"
        for request in call["body"]["requests"]
    ]
    for request in batch_requests:
        if "repeatCell" in request:
            cell_range = request["repeatCell"]["range"]
            assert cell_range.get("endRowIndex", 1) > cell_range.get("startRowIndex", 0)


def test_output_config_rejects_invalid_backend_paths() -> None:
    import pytest

    with pytest.raises(ValueError, match="output.backend"):
        validate_output_config({"output": {"backend": "csv"}})
    with pytest.raises(ValueError, match="output.xlsx.directory"):
        validate_output_config(
            {"output": {"backend": "xlsx", "xlsx": {"directory": []}}}
        )
    with pytest.raises(ValueError, match="credentials_file"):
        validate_output_config(
            {
                "output": {
                    "backend": "google_sheets",
                    "google_sheets": {"credentials_file": ""},
                }
            }
        )
